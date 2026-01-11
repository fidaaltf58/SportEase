"""
Script d'export des résultats de tests SportEase vers Excel
Auteur: Équipe SportEase
Date: Janvier 2025
Usage: python generate_report.py
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, Reference
from datetime import datetime
import os

# ===== CONFIGURATION =====
PROJECT_NAME = "SportEase"
VERSION = "v1.0-Sprint6"
OUTPUT_FILE = "SportEase_Test_Report.xlsx"

# ===== DONNÉES DE TEST =====
test_results = [
    # Authentification (5 tests)
    {'test_case_id': 'TC-001', 'test_name': 'Inscription Utilisateur', 'component': 'Authentification', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '5 min', 'comments': 'Test réussi'},
    {'test_case_id': 'TC-002', 'test_name': 'Inscription Email Existant', 'component': 'Authentification', 'priority': 'HIGH', 'status': 'FAILED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '3 min', 'comments': 'BUG-002: Message générique'},
    {'test_case_id': 'TC-003', 'test_name': 'Connexion Réussie', 'component': 'Authentification', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-11', 'duration': '2 min', 'comments': ''},
    {'test_case_id': 'TC-004', 'test_name': 'Connexion Échouée', 'component': 'Authentification', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-11', 'duration': '2 min', 'comments': ''},
    {'test_case_id': 'TC-005', 'test_name': 'Modification Profil', 'component': 'Authentification', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '4 min', 'comments': ''},
    
    # Gestion Terrains (8 tests)
    {'test_case_id': 'TC-006', 'test_name': 'Création Terrain', 'component': 'Gestion Terrains', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '8 min', 'comments': 'Upload image OK'},
    {'test_case_id': 'TC-007', 'test_name': 'Modification Terrain', 'component': 'Gestion Terrains', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-008', 'test_name': 'Archivage Terrain', 'component': 'Gestion Terrains', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '3 min', 'comments': ''},
    {'test_case_id': 'TC-009', 'test_name': 'Upload Image', 'component': 'Gestion Terrains', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '10 min', 'comments': 'Test JPG, PNG, GIF'},
    {'test_case_id': 'TC-010', 'test_name': 'Validation Prix', 'component': 'Gestion Terrains', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '2 min', 'comments': ''},
    {'test_case_id': 'TC-011', 'test_name': 'Validation Horaires', 'component': 'Gestion Terrains', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '3 min', 'comments': ''},
    {'test_case_id': 'TC-012', 'test_name': 'Liste Admin', 'component': 'Gestion Terrains', 'priority': 'LOW', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '4 min', 'comments': ''},
    {'test_case_id': 'TC-013', 'test_name': 'Statistiques Terrains', 'component': 'Gestion Terrains', 'priority': 'LOW', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-11', 'duration': '3 min', 'comments': ''},
    
    # Consultation Terrains (6 tests)
    {'test_case_id': 'TC-014', 'test_name': 'Liste Terrains', 'component': 'Consultation Terrains', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-015', 'test_name': 'Recherche Sport', 'component': 'Consultation Terrains', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '4 min', 'comments': ''},
    {'test_case_id': 'TC-016', 'test_name': 'Recherche Multi-critères', 'component': 'Consultation Terrains', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '6 min', 'comments': ''},
    {'test_case_id': 'TC-017', 'test_name': 'Détails Terrain', 'component': 'Consultation Terrains', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-018', 'test_name': 'Disponibilités Temps Réel', 'component': 'Consultation Terrains', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '7 min', 'comments': ''},
    {'test_case_id': 'TC-019', 'test_name': 'Filtres et Tri', 'component': 'Consultation Terrains', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '8 min', 'comments': ''},
    
    # Réservations (12 tests)
    {'test_case_id': 'TC-020', 'test_name': 'Création Réservation', 'component': 'Réservations', 'priority': 'CRITICAL', 'status': 'FAILED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '10 min', 'comments': 'BUG-001: Prix incorrect'},
    {'test_case_id': 'TC-021', 'test_name': 'Conflit Créneau', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'FAILED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '15 min', 'comments': 'BUG-003: Double réservation'},
    {'test_case_id': 'TC-022', 'test_name': 'Limite 3 Réservations', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-023', 'test_name': 'Validation Durée', 'component': 'Réservations', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '3 min', 'comments': ''},
    {'test_case_id': 'TC-024', 'test_name': 'Réservation Passé', 'component': 'Réservations', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '3 min', 'comments': ''},
    {'test_case_id': 'TC-025', 'test_name': 'Réservation >30j', 'component': 'Réservations', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '3 min', 'comments': ''},
    {'test_case_id': 'TC-026', 'test_name': 'Annulation >24h', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-027', 'test_name': 'Annulation <24h', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '4 min', 'comments': ''},
    {'test_case_id': 'TC-028', 'test_name': 'Confirmation Admin', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '4 min', 'comments': ''},
    {'test_case_id': 'TC-029', 'test_name': 'Refus Admin', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-030', 'test_name': 'Calcul Prix', 'component': 'Réservations', 'priority': 'HIGH', 'status': 'BLOCKED', 'executed_by': 'Mohamed Ali', 'execution_date': '2025-01-12', 'duration': '2 min', 'comments': 'Dépend BUG-001'},
    {'test_case_id': 'TC-031', 'test_name': 'Mes Réservations', 'component': 'Réservations', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    
    # Dashboards (4 tests)
    {'test_case_id': 'TC-032', 'test_name': 'Dashboard User', 'component': 'Dashboards', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '6 min', 'comments': ''},
    {'test_case_id': 'TC-033', 'test_name': 'Dashboard Admin', 'component': 'Dashboards', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '7 min', 'comments': ''},
    {'test_case_id': 'TC-034', 'test_name': 'Statistiques Admin', 'component': 'Dashboards', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-035', 'test_name': 'Graphiques Chart.js', 'component': 'Dashboards', 'priority': 'LOW', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-12', 'duration': '3 min', 'comments': ''},
    
    # Sécurité (10 tests)
    {'test_case_id': 'TC-036', 'test_name': 'SQL Injection', 'component': 'Sécurité', 'priority': 'CRITICAL', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '15 min', 'comments': 'Protection efficace'},
    {'test_case_id': 'TC-037', 'test_name': 'XSS Protection', 'component': 'Sécurité', 'priority': 'CRITICAL', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '12 min', 'comments': 'Scripts échappés'},
    {'test_case_id': 'TC-038', 'test_name': 'CSRF Token', 'component': 'Sécurité', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '5 min', 'comments': ''},
    {'test_case_id': 'TC-039', 'test_name': 'Accès Non Autorisé', 'component': 'Sécurité', 'priority': 'CRITICAL', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '8 min', 'comments': ''},
    {'test_case_id': 'TC-040', 'test_name': 'Mots de Passe Hachés', 'component': 'Sécurité', 'priority': 'CRITICAL', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '10 min', 'comments': 'BCrypt utilisé'},
    {'test_case_id': 'TC-041', 'test_name': 'Session Timeout', 'component': 'Sécurité', 'priority': 'HIGH', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '15 min', 'comments': '30 min timeout'},
    {'test_case_id': 'TC-042', 'test_name': 'Upload Fichiers', 'component': 'Sécurité', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '10 min', 'comments': 'Validation MIME'},
    {'test_case_id': 'TC-043', 'test_name': 'Headers Sécurité', 'component': 'Sécurité', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '5 min', 'comments': 'X-Frame-Options présent'},
    {'test_case_id': 'TC-044', 'test_name': 'Logs Sécurité', 'component': 'Sécurité', 'priority': 'LOW', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '3 min', 'comments': ''},
    {'test_case_id': 'TC-045', 'test_name': 'HTTPS Redirect', 'component': 'Sécurité', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '2 min', 'comments': ''},
    
    # Performance (3 tests)
    {'test_case_id': 'TC-046', 'test_name': 'Chargement Accueil', 'component': 'Performance', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '5 min', 'comments': '1.2s average'},
    {'test_case_id': 'TC-047', 'test_name': 'Recherche <1s', 'component': 'Performance', 'priority': 'MEDIUM', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '5 min', 'comments': '0.8s average'},
    {'test_case_id': 'TC-048', 'test_name': 'Test Performance', 'component': 'Performance', 'priority': 'LOW', 'status': 'PASSED', 'executed_by': 'Ahmed Ben Salem', 'execution_date': '2025-01-13', 'duration': '20 min', 'comments': 'Apache Bench: 97 req/sec'},
]

# Bugs identifiés
bugs = [
    {'id': 'BUG-001', 'severity': 'Critique', 'test_case': 'TC-020', 'component': 'Réservations', 
     'description': 'Calcul incorrect du prix pour réservations multi-heures. 2h à 50€/h affiche 75€ au lieu de 100€.',
     'status': 'Ouvert', 'assigned': 'Dev Team', 'date': '2025-01-12'},
     
    {'id': 'BUG-002', 'severity': 'Majeur', 'test_case': 'TC-002', 'component': 'Authentification',
     'description': 'Message générique "Une erreur est survenue" au lieu de message spécifique pour email existant.',
     'status': 'Ouvert', 'assigned': 'Dev Team', 'date': '2025-01-11'},
     
    {'id': 'BUG-003', 'severity': 'Majeur', 'test_case': 'TC-021', 'component': 'Réservations',
     'description': 'Deux utilisateurs peuvent réserver le même créneau horaire simultanément.',
     'status': 'Ouvert', 'assigned': 'Dev Team', 'date': '2025-01-12'},
]

# ===== FONCTIONS =====

def create_summary_sheet(wb):
    """Créer la feuille de résumé"""
    ws = wb.create_sheet('Résumé', 0)
    
    # En-tête
    ws['A1'] = f'RAPPORT DE TESTS - {PROJECT_NAME}'
    ws['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:E1')
    ws.row_dimensions[1].height = 30
    
    # Informations projet
    ws['A3'] = 'Projet:'
    ws['B3'] = PROJECT_NAME
    ws['A4'] = 'Version:'
    ws['B4'] = VERSION
    ws['A5'] = 'Date du rapport:'
    ws['B5'] = datetime.now().strftime('%d/%m/%Y %H:%M')
    ws['A6'] = 'Testeur(s):'
    ws['B6'] = 'Ahmed Ben Salem, Mohamed Ali'
    
    # Style labels
    for row in range(3, 7):
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    
    # Statistiques globales
    total = len(test_results)
    passed = sum(1 for t in test_results if t['status'] == 'PASSED')
    failed = sum(1 for t in test_results if t['status'] == 'FAILED')
    blocked = sum(1 for t in test_results if t['status'] == 'BLOCKED')
    rate = (passed / total * 100) if total > 0 else 0
    
    ws['A8'] = 'STATISTIQUES GLOBALES'
    ws['A8'].font = Font(size=12, bold=True, color='FFFFFF')
    ws['A8'].fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
    ws.merge_cells('A8:E8')
    
    stats = [
        ['Métrique', 'Valeur'],
        ['Total de tests', total],
        ['Tests réussis', passed],
        ['Tests échoués', failed],
        ['Tests bloqués', blocked],
        ['Taux de réussite', f'{rate:.2f}%'],
    ]
    
    for i, (label, value) in enumerate(stats, start=9):
        ws[f'A{i}'] = label
        ws[f'B{i}'] = value
        for col in ['A', 'B']:
            ws[f'{col}{i}'].border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                           top=Side(style='thin'), bottom=Side(style='thin'))
        if i == 9:  # Header
            for col in ['A', 'B']:
                ws[f'{col}{i}'].font = Font(bold=True)
                ws[f'{col}{i}'].fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    
    # Graphique camembert
    pie = PieChart()
    labels = Reference(ws, min_col=1, min_row=11, max_row=13)
    data = Reference(ws, min_col=2, min_row=10, max_row=13)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.title = "Répartition des Résultats"
    ws.add_chart(pie, "D3")
    
    # Résultats par composant
    ws['A16'] = 'RÉSULTATS PAR COMPOSANT'
    ws['A16'].font = Font(size=12, bold=True, color='FFFFFF')
    ws['A16'].fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
    ws.merge_cells('A16:E16')
    
    headers = ['Composant', 'Total', 'Réussis', 'Échoués', 'Taux %']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=17, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                            top=Side(style='thin'), bottom=Side(style='thin'))
    
    # Calcul par composant
    components = {}
    for test in test_results:
        comp = test['component']
        if comp not in components:
            components[comp] = {'total': 0, 'passed': 0, 'failed': 0}
        components[comp]['total'] += 1
        if test['status'] == 'PASSED':
            components[comp]['passed'] += 1
        elif test['status'] == 'FAILED':
            components[comp]['failed'] += 1
    
    row = 18
    for comp, stats in components.items():
        rate = (stats['passed'] / stats['total'] * 100) if stats['total'] > 0 else 0
        data = [comp, stats['total'], stats['passed'], stats['failed'], f'{rate:.2f}%']
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                               top=Side(style='thin'), bottom=Side(style='thin'))
        row += 1
    
    # Ajuster largeurs
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15

def create_detailed_results_sheet(wb):
    """Créer la feuille des résultats détaillés"""
    ws = wb.create_sheet('Résultats Détaillés')
    
    headers = ['ID', 'Nom du Test', 'Composant', 'Priorité', 'Statut', 
               'Exécuté par', 'Date', 'Durée', 'Commentaires']
    
    # En-tête
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                           top=Side(style='thin'), bottom=Side(style='thin'))
    
    # Données
    for idx, test in enumerate(test_results, start=2):
        data = [
            test['test_case_id'],
            test['test_name'],
            test['component'],
            test['priority'],
            test['status'],
            test['executed_by'],
            test['execution_date'],
            test['duration'],
            test['comments']
        ]
        
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=idx, column=col, value=value)
            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                               top=Side(style='thin'), bottom=Side(style='thin'))
            
            # Couleur selon statut
            if col == 5:  # Colonne Statut
                if value == 'PASSED':
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                    cell.font = Font(color='006100', bold=True)
                elif value == 'FAILED':
                    cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                    cell.font = Font(color='9C0006', bold=True)
                elif value == 'BLOCKED':
                    cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                    cell.font = Font(color='9C6500', bold=True)
            
            # Couleur selon priorité
            elif col == 4:  # Colonne Priorité
                if value == 'CRITICAL':
                    cell.font = Font(color='FF0000', bold=True)
                elif value == 'HIGH':
                    cell.font = Font(color='FF6B00', bold=True)
                elif value == 'MEDIUM':
                    cell.font = Font(color='0070C0', bold=True)
                elif value == 'LOW':
                    cell.font = Font(color='808080')
    
    # Ajuster largeurs
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 10
    ws.column_dim